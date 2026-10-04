import docx
import os

doc = docx.Document('23f1000204ds.study.iitm.ac.in - Mid Term.docx')

page_breaks = 0
for p in doc.paragraphs:
    xml = p._element.xml
    if 'w:br' in xml and 'w:type="page"' in xml:
        page_breaks += 1

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")
print(f"Explicit Page Breaks: {page_breaks}")

# Try win32com to get exact page count if Word is installed
try:
    import win32com.client
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_path = os.path.abspath('23f1000204ds.study.iitm.ac.in - Mid Term.docx')
    wdoc = word.Documents.Open(doc_path)
    pages = wdoc.ComputeStatistics(2) # 2 = wdStatisticPages
    wdoc.Close(False)
    word.Quit()
    print(f"EXACT MS WORD PAGE COUNT: {pages}")
except Exception as e:
    print(f"Win32com check info: {e}")
