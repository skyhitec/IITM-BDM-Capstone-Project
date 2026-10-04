from docx import Document

doc = Document('23f1000204ds.study.iitm.ac.in - Mid Term.docx')
print(f'Paragraphs: {len(doc.paragraphs)}')
print(f'Tables: {len(doc.tables)}')

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if text:
        safe = text.encode('ascii','replace').decode('ascii')
        style = para.style.name if para.style else 'None'
        print(f'[P{i:02d}] [{style}] {safe[:200]}')

print('\n=== TABLES ===')
for t, table in enumerate(doc.tables):
    print(f'\nTable {t+1}: {len(table.rows)}x{len(table.columns)}')
    for r, row in enumerate(table.rows):
        if r < 3:
            vals = [cell.text.strip()[:30].encode('ascii','replace').decode() for cell in row.cells]
            print(f'  {vals}')
