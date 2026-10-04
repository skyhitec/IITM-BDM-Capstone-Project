from docx import Document
import os

doc_path = 'c:\\Users\\HP\\Downloads\\BDM\\Mid term.docx'
if not os.path.exists(doc_path):
    print("File not found:", doc_path)
else:
    doc = Document(doc_path)
    print(f"File Size: {os.path.getsize(doc_path)/1024:.1f} KB")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    
    images = [rel for rel in doc.part.rels.values() if "image" in rel.reltype]
    print(f"Total Images/Figures: {len(images)}")
    
    print("\n=== PARAGRAPHS & HEADINGS ===")
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        if t:
            style = p.style.name if p.style else 'Normal'
            safe = t.encode('ascii', 'replace').decode()
            print(f"[P{i:03d}] [{style}] {safe[:120]}")
            
    print("\n=== TABLES ===")
    for t, table in enumerate(doc.tables):
        print(f"Table {t+1}: {len(table.rows)}x{len(table.columns)}")
        for r, row in enumerate(table.rows[:3]):
            vals = [c.text.strip()[:30].encode('ascii','replace').decode() for c in row.cells]
            print(f"  Row {r}: {vals}")
