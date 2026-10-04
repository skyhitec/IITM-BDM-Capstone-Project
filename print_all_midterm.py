from docx import Document
import os

doc = Document('c:\\Users\\HP\\Downloads\\BDM\\Mid term.docx')

print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t:
        style = p.style.name if p.style else 'None'
        safe = t.encode('ascii','replace').decode()
        print(f"[{i:03d}] [{style}] {safe}")
